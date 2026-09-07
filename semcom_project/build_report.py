"""Build the final academic DOCX from 技术报告.md with deterministic styling."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parent
SOURCE_PATH = PROJECT_DIR / "技术报告.md"
OUTPUT_PATH = PROJECT_DIR.parent / "output" / "docx" / "语义通信课程设计技术报告.docx"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "DCE6F1"
LIGHT_GRAY = "F2F4F7"
MUTED = "666666"
GOLD = "B08D32"
CONTENT_WIDTH_DXA = 9025  # A4 width minus 2.54 cm margins on both sides.


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, east_asia="宋体", ascii_font="Times New Roman", size=10.5, bold=None, color=None, italic=None):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def strip_inline_markdown(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("**", "")
    return text


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(21)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5

    styles = [
        ("Heading 1", "宋体", 16, NAVY, 14, 8),
        ("Heading 2", "宋体", 13, BLUE, 12, 6),
        ("Heading 3", "宋体", 11.5, DARK_BLUE, 8, 4),
    ]
    for name, font, size, color, before, after in styles:
        style = document.styles[name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:ascii"), "SimSun")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "SimSun")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True

    caption = document.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0

    code = document.styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    code.font.size = Pt(9)
    code.paragraph_format.first_line_indent = Pt(0)
    code.paragraph_format.left_indent = Cm(0.6)
    code.paragraph_format.right_indent = Cm(0.4)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(3)
    code.paragraph_format.line_spacing = 1.0


def set_page_geometry(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run("第 ")
    set_run_font(run, east_asia="宋体", size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, east_asia="宋体", size=9, color=MUTED)


def configure_headers_footers(document: Document) -> None:
    section = document.sections[0]
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("现代通信技术实践  |  Deep JSCC 图像语义通信")
    set_run_font(run, east_asia="等线", ascii_font="Arial", size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(document: Document) -> None:
    for _ in range(3):
        document.add_paragraph()
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.first_line_indent = Pt(0)
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("现代通信技术实践 · 课程设计技术报告")
    set_run_font(run, east_asia="宋体", ascii_font="SimSun", size=12, bold=True, color=GOLD)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(16)
    run = title.add_run("基于深度学习的图像语义通信系统\n设计与实践")
    set_run_font(run, east_asia="宋体", ascii_font="SimSun", size=25, bold=True, color=NAVY)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    subtitle.paragraph_format.space_after = Pt(56)
    run = subtitle.add_run("Deep JSCC 复现、模型改进与鲁棒性分析")
    set_run_font(run, east_asia="宋体", ascii_font="SimSun", size=16, bold=False, color=BLUE)

    metadata = document.add_table(rows=5, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.autofit = False
    labels = ["课程", "方向", "小组成员 / 学号", "指导教师", "完成日期"]
    values = [
        "现代通信技术实践",
        "基于深度学习的语义通信系统设计与实践",
        "____________________________",
        "____________________________",
        "2026 年 8 月",
    ]
    widths = [2500, 5000]
    for row_index, row in enumerate(metadata.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = Cm(widths[column_index] / 567)
            set_cell_margins(cell, top=130, bottom=130, start=160, end=160)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            text = labels[row_index] if column_index == 0 else values[row_index]
            run = paragraph.add_run(text)
            set_run_font(run, east_asia="宋体", size=11, bold=column_index == 0, color=NAVY if column_index == 0 else None)
            if column_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
    set_table_geometry(metadata, [2500, 5000], table_width=7500, indent=760)
    document.add_page_break()


def set_table_geometry(table, widths: list[int], table_width=CONTENT_WIDTH_DXA, indent=120) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(table_width))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    table.autofit = False
    weights = []
    for column in range(column_count):
        longest = max(len(row[column]) for row in rows)
        weights.append(max(5.0, min(24.0, float(longest))))
    if column_count >= 4:
        weights[0] = max(weights[0], 13.0)
    total = sum(weights)
    widths = [max(760, int(CONTENT_WIDTH_DXA * weight / total)) for weight in weights]
    difference = CONTENT_WIDTH_DXA - sum(widths)
    widths[-1] += difference
    set_table_geometry(table, widths)

    font_size = 8.0 if column_count >= 7 else 8.5 if column_count >= 6 else 9.0 if column_count >= 5 else 9.5
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 and row_index > 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(strip_inline_markdown(value))
            set_run_font(run, east_asia="黑体" if row_index == 0 else "宋体", size=font_size, bold=row_index == 0)
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_tr_pr.append(repeat)
    after = document.add_paragraph()
    after.paragraph_format.first_line_indent = Pt(0)
    after.paragraph_format.space_after = Pt(2)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_toc(document: Document) -> None:
    heading = document.add_paragraph("目录", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.page_break_before = False
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在打开文档或导出 PDF 时自动更新。"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])
    document.add_page_break()


def parse_markdown(document: Document, lines: list[str]) -> None:
    index = 0
    in_code = False
    code_lines: list[str] = []
    toc_added = False
    while index < len(lines):
        raw = lines[index].rstrip()
        line = raw.strip()
        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("课程：") or line.startswith("方向：") or line.startswith("小组成员：") or line.startswith("指导教师：") or line.startswith("完成日期："):
            index += 1
            continue
        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                shade_paragraph(paragraph, LIGHT_GRAY)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, east_asia="等线", ascii_font="Consolas", size=9)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            text = line[3:].strip()
            if text == "1 引言" and not toc_added:
                add_toc(document)
                toc_added = True
            paragraph = document.add_paragraph(text, style="Heading 1")
            if text in {"附录 A 代码结构与复现命令", "附录 B 课程要求覆盖对照"}:
                paragraph.paragraph_format.page_break_before = True
            index += 1
            continue
        if line.startswith("### "):
            subsection = line[4:].strip()
            paragraph = document.add_paragraph(subsection, style="Heading 2")
            index += 1
            continue
        image_match = re.match(r"!\[(.+?)\]\((.+?)\)", line)
        if image_match:
            caption, relative_path = image_match.groups()
            path = PROJECT_DIR / relative_path
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run()
            shape = run.add_picture(str(path), width=Cm(15.2))
            shape._inline.docPr.set("descr", caption)
            shape._inline.docPr.set("title", caption)
            document.add_paragraph(caption, style="Caption")
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    continue
                parsed.append(cells)
            add_table(document, parsed)
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if candidate.startswith(("#", "|", "```", "![")):
                break
            paragraph_lines.append(candidate)
            index += 1
        text = strip_inline_markdown(" ".join(paragraph_lines))
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.widow_control = True
        if text.startswith("关键词："):
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(10)
            label, value = text.split("：", 1)
            label_run = paragraph.add_run(label + "：")
            set_run_font(label_run, east_asia="黑体", size=10.5, bold=True)
            value_run = paragraph.add_run(value)
            set_run_font(value_run, east_asia="宋体", size=10.5)
        elif len(text) < 90 and any(token in text for token in ("arg min", "log10", "CN(", "k/n =", "E[z", "z =", "y =")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            run = paragraph.add_run(text)
            set_run_font(run, east_asia="宋体", size=10.5, italic=True)
        elif re.match(r"^\[\d+\]", text):
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.hanging_indent = Cm(0.6)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(text)
            set_run_font(run, east_asia="宋体", size=9.5)
        else:
            run = paragraph.add_run(text)
            set_run_font(run, east_asia="宋体", size=10.5)


def set_update_fields_on_open(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    document = Document()
    section = document.sections[0]
    set_page_geometry(section)
    configure_styles(document)
    configure_headers_footers(document)
    document.core_properties.title = "基于深度学习的图像语义通信系统设计与实践"
    document.core_properties.subject = "现代通信技术实践课程设计"
    document.core_properties.author = "课程设计小组"
    document.core_properties.keywords = "语义通信; Deep JSCC; AWGN; 图像重建"
    add_cover(document)
    lines = SOURCE_PATH.read_text(encoding="utf-8").splitlines()
    parse_markdown(document, lines)
    set_update_fields_on_open(document)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
